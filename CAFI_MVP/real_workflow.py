"""Orquestación real, independiente del simulador y con diario de impresión."""
import io
import json
import os
import threading
import uuid
from pathlib import Path

from openpyxl.styles import Font, PatternFill
from config import Rules
from excel_manager import ALIASES, open_excel
from journal import Journal, worker_lock
from windows_rpa import RpaError, StopRequested, WindowsDriver, load_profile


class RealJournal(Journal):
    """Usar siempre data/operacion; jamás importar estados del simulador."""
    def __init__(self, root):
        super().__init__(root)
        with self.db() as db:
            db.execute('''CREATE TABLE IF NOT EXISTS outputs (
                run_id TEXT NOT NULL, row_no INTEGER NOT NULL, phase TEXT NOT NULL,
                payload TEXT NOT NULL DEFAULT '{}', PRIMARY KEY(run_id,row_no))''')

    def phase(self, run_id, row_no, phase, payload=None):
        with self.db() as db:
            db.execute('''INSERT INTO outputs VALUES(?,?,?,?) ON CONFLICT(run_id,row_no)
                DO UPDATE SET phase=excluded.phase, payload=CASE WHEN excluded.payload='{}'
                THEN outputs.payload ELSE excluded.payload END''',
                (run_id, row_no, phase, json.dumps(payload or {}, ensure_ascii=False)))

    def outputs(self, run_id):
        with self.db() as db:
            return {r['row_no']: {'phase': r['phase'], **json.loads(r['payload'])}
                    for r in db.execute('SELECT * FROM outputs WHERE run_id=?', (run_id,))}

    def check_duplicates(self, run):
        with self.db() as db:
            for row in run['rows']:
                prior = db.execute('''SELECT run_id FROM rows WHERE fingerprint=? AND run_id<>?
                    AND state IN ('PROCESADO','EN_PROCESO','REQUIERE_REVISION') LIMIT 1''',
                    (row['fingerprint'], run['id'])).fetchone()
                if prior:
                    raise RpaError(f'Fila {row["row"]}: existe un intento previo con el mismo contenido. Concilie antes de continuar.')

    def check_transaction(self, transaction_id):
        with self.db() as db:
            if db.execute('SELECT 1 FROM rows WHERE transaction_id=? LIMIT 1', (transaction_id,)).fetchone():
                raise RpaError('El reporte muestra un código de trámite ya registrado. Revise si quedó un reporte anterior abierto.')

    def export(self, run_id):
        run = self.get(run_id)
        folder = self.root / run_id
        wb = open_excel((folder / 'origen.xlsx').read_bytes())
        ws = wb[run['sheet']]
        columns = {ALIASES.get(str(c.value).strip(), str(c.value)): c.column for c in ws[1] if c.value}
        keys = ['COD_TRAMITE', 'GPA', 'CREDITOS_APROBADOS', 'ESTADO', 'MENSAJE_LOG', 'IMPRESION_CAFI', 'MODO_CAFI', 'RUN_ID_CAFI']
        for key in keys:
            if key not in columns:
                columns[key] = ws.max_column + 1
                cell = ws.cell(1, columns[key], key)
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='8A1538')
                ws.column_dimensions[cell.column_letter].width = 26
        outputs = self.outputs(run_id)
        for row in run['rows']:
            out = outputs.get(row['row'], {})
            values = [out.get('transaction_id', row.get('transaction_id') or ''), out.get('gpa', ''),
                      out.get('credits', ''), row['state'], row['message'], out.get('phase', 'PENDIENTE'),
                      'RPA WINDOWS', run_id]
            for key, value in zip(keys, values):
                cell = ws.cell(row['row'], columns[key])
                cell.value = str(value)
                cell.data_type = 's'  # Texto observado nunca se interpreta como fórmula.
        temporary = folder / (uuid.uuid4().hex + '.tmp')
        try:
            stream = io.BytesIO()
            wb.save(stream)
            with open(temporary, 'xb') as file:
                file.write(stream.getvalue())
                file.flush()
                os.fsync(file.fileno())
            os.replace(temporary, folder / 'resultados_cafi.xlsx')
        finally:
            wb.close()
            temporary.unlink(missing_ok=True)
        return folder / 'resultados_cafi.xlsx'


def execute_rows(journal, run_id, driver, stop_event, progress=lambda message: None, limit=None,
                 confirm_results=None):
    """El llamador mantiene worker_lock durante toda la conexión y ejecución."""
    run = journal.get(run_id)
    if any(r['issues'] for r in run['rows']):
        raise RpaError('El Excel contiene incidencias. Corríjalas antes de iniciar.')
    if any(r['state'] in ('REQUIERE_REVISION', 'EN_PROCESO') for r in run['rows']):
        raise RpaError('Hay un guardado incierto. Se requiere conciliación en CAFI.')
    outputs = journal.outputs(run_id)
    if any(r['state'] == 'PROCESADO' and outputs.get(r['row'], {}).get('phase') != 'COMPLETO' for r in run['rows']):
        raise RpaError('Hay un trámite guardado cuya impresión o limpieza no terminó. Revíselo en CAFI; no se repetirá.')
    journal.check_duplicates(run)
    journal.export(run_id)
    journal.status(run_id, 'EN_EJECUCION', 'RPA en ejecución sobre la ventana de CAFI seleccionada.')
    current = None
    sent = False
    success_observed = False
    count = 0
    code = 0
    try:
        for row in run['rows']:
            if row['state'] == 'PROCESADO':
                continue
            if stop_event.is_set() or (limit is not None and count >= limit):
                code = 3
                break
            current, sent, success_observed = row, False, False
            count += 1
            n = row['row']
            progress(f'Fila {n} · Completando el formulario y verificando el destinatario…')
            journal.transition(run_id, n, 'EN_PROCESO', 'PREPARACION', 'Preparando formulario CAFI.')
            driver.prepare(row)
            journal.transition(run_id, n, 'VALIDADO', 'LECTURA', 'Destinatario y contenido leídos de vuelta en CAFI.')
            if stop_event.is_set():
                raise StopRequested('Detenido antes de guardar. El formulario puede conservar los datos; revíselo antes de retomar.')
            journal.transition(run_id, n, 'EN_PROCESO', 'INTENCION_GUARDAR', 'Intención persistida; no reenviar si el resultado es incierto.')
            sent = True
            progress(f'Fila {n} · Guardando y esperando la confirmación de CAFI…')
            driver.save()
            success_observed = True
            journal.transition(run_id, n, 'PROCESADO', 'EXITO', 'CAFI confirmó el guardado. Impresión pendiente.')
            journal.phase(run_id, n, 'GUARDADO')
            journal.export(run_id)
            progress(f'Fila {n} · Preparando la impresión…')
            driver.open_print()
            journal.phase(run_id, n, 'INTENCION_IMPRIMIR')
            driver.print_report()
            journal.phase(run_id, n, 'REPORTE_ABIERTO')
            result = driver.results()
            journal.check_transaction(result['transaction_id'])
            if confirm_results:
                journal.phase(run_id, n, 'PENDIENTE_CONFIRMACION')
                progress(f'Fila {n} · Reporte listo. Confirma los resultados antes de actualizar el Excel.')
                if not confirm_results(row, result):
                    raise RpaError('La usuaria no confirmó los resultados impresos. La fila queda bloqueada para conciliación y no se actualizará el Excel.')
            journal.phase(run_id, n, 'RESULTADOS_LEIDOS', result)
            journal.transition(run_id, n, 'PROCESADO', 'RESULTADOS', 'Reporte generado; código, GPA y créditos leídos.', result['transaction_id'])
            journal.export(run_id)
            progress(f'Fila {n} · Resultados registrados. Preparando el siguiente trámite…')
            driver.reset()
            journal.phase(run_id, n, 'COMPLETO')
            journal.transition(run_id, n, 'PROCESADO', 'COMPLETO', 'Guardado, reporte y limpieza del formulario confirmados.')
            journal.export(run_id)
            current = None
        journal.status(run_id, 'COMPLETADO' if code == 0 else 'DETENIDO',
                       'Todos los trámites del lote terminaron.' if code == 0 else 'Detenido antes del siguiente registro.')
    except (Exception, KeyboardInterrupt) as exc:
        code = 3
        if current and not success_observed:
            journal.transition(run_id, current['row'], 'REQUIERE_REVISION' if sent else 'ERROR',
                               'DETENCION', 'Guardado incierto. Concilie en CAFI; no se reintentará.' if sent else 'Detenido antes de Guardar; revise el formulario antes de retomar.')
        # Los errores externos podrían incluir contraseña o texto de una celda: no se serializan.
        message = str(exc) if isinstance(exc, RpaError) else 'Fallo de conexión o escritura. Revise CAFI y la bitácora; no se repitió el guardado.'
        if current and success_observed:
            message = 'CAFI confirmó el guardado; la continuación requiere revisión. No vuelva a cargar esta fila. ' + message
        journal.status(run_id, 'DETENIDO', message)
        progress(message)
    finally:
        try:
            journal.export(run_id)
        except OSError:
            journal.status(run_id, 'DETENIDO', 'Cierre resultados_cafi.xlsx y vuelva a exportar. La bitácora conserva el avance.')
            progress('No se pudo actualizar el Excel de salida. Cierre ese archivo y use Exportar resultados.')
            code = 3
    return code


def start_workflow(attached, workbook, profile_path, data_root, *, username='', password='',
                   authenticated=False, rules=Rules(), stop_event=None, progress=lambda message: None, limit=None,
                   confirm_results=None):
    stop_event = stop_event or threading.Event()
    profile = load_profile(profile_path, login=not authenticated)
    if not authenticated and (not username.strip() or not password):
        raise RpaError('Ingrese el usuario y la contraseña, o marque que CAFI ya tiene sesión iniciada.')
    if Path(workbook).suffix.lower() != '.xlsx':
        raise RpaError('Seleccione un archivo .xlsx.')
    journal = RealJournal(data_root)
    with worker_lock(journal.root):
        journal.recover()
        run = journal.create(Path(workbook).read_bytes(), rules)
        if any(r['issues'] for r in run['rows']):
            raise RpaError('Hay incidencias en el Excel; use Validar Excel para verlas.')
        # Comprobar estados antes de iniciar sesión o tocar el formulario.
        journal.check_duplicates(run)
        if any(r['state'] in ('REQUIERE_REVISION', 'EN_PROCESO') for r in run['rows']):
            raise RpaError('Este lote tiene intentos inciertos. Concilie en CAFI antes de continuar.')
        outputs = journal.outputs(run['id'])
        if any(r['state'] == 'PROCESADO' and outputs.get(r['row'], {}).get('phase') != 'COMPLETO' for r in run['rows']):
            raise RpaError('Hay un trámite guardado con continuación incompleta. Revise su impresión en CAFI; no se repetirá.')
        if all(r['state'] == 'PROCESADO' for r in run['rows']):
            return {'code': 0, 'run_id': run['id'], 'output': str(journal.export(run['id'])),
                    'message': 'Este lote ya terminó. Se recuperaron sus resultados sin volver a operar CAFI.'}
        driver = WindowsDriver(attached, profile, stop_event)
        try:
            if not authenticated:
                progress('Iniciando sesión en CAFI…')
                driver.login(username, password)
        finally:
            username = password = ''  # Sin almacenamiento persistente; Python no garantiza borrado físico de memoria.
        code = execute_rows(journal, run['id'], driver, stop_event, progress, limit,
                            confirm_results=confirm_results)
        return {'code': code, 'run_id': run['id'], 'output': str(journal.root / run['id'] / 'resultados_cafi.xlsx'),
                'message': journal.metadata(run['id'])['message']}
