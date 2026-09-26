"""Entrada determinista. Nunca escribe sobre el archivo original."""
import hashlib
import io
import json
import re
import zipfile
from dataclasses import asdict, dataclass, field

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from config import Rules

ALIASES = {
    'COD. ALUMNO': 'COD_ALUMNO', 'COD_ALUMNO': 'COD_ALUMNO',
    'COD_ALUMNO_O_FACULTAD': 'COD_ALUMNO',
    'NOMBRE ESTUDIANTE': 'NOMBRE_ESTUDIANTE', 'NOMBRE_ESTUDIANTE': 'NOMBRE_ESTUDIANTE',
    'TITULO': 'TITULO', 'DETALLE': 'DETALLE', 'PRIORIDAD': 'PRIORIDAD',
    'CÓDIGO:': 'COD_TRAMITE', 'COD_TRAMITE': 'COD_TRAMITE',
    'GPA:': 'GPA', 'GPA': 'GPA', 'CRED.APROB:': 'CREDITOS_APROBADOS',
    'CREDITOS_APROBADOS': 'CREDITOS_APROBADOS', 'ESTADO': 'ESTADO',
    'MENSAJE_LOG': 'MENSAJE_LOG',
}
REQUIRED = {'COD_ALUMNO', 'TITULO', 'DETALLE'}

class InputError(ValueError):
    pass

@dataclass
class Row:
    row: int
    code: str = ''
    name: str = ''
    title: str = ''
    detail: str = ''
    priority: str = 'Media'
    issues: list = field(default_factory=list)
    notes: list = field(default_factory=list)
    fingerprint: str = ''

def digest(content):
    return hashlib.sha256(content).hexdigest()

def open_excel(content):
    if len(content) > 10 * 1024 * 1024:
        raise InputError('El archivo supera 10 MB. Divida el lote.')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 50 * 1024 * 1024:
                raise InputError('El libro expandido supera el límite de 50 MB.')
        return load_workbook(io.BytesIO(content), data_only=False, keep_links=False)
    except InputError:
        raise
    except Exception as exc:
        raise InputError('No se pudo leer el libro. Use un archivo .xlsx válido, sin contraseña.') from exc

def code_value(cell, length):
    value = cell.value
    if cell.data_type == 'f':
        raise InputError('Código con fórmula: conviértalo en texto revisado.')
    if isinstance(value, bool) or value is None:
        raise InputError('Código vacío o inválido.')
    if isinstance(value, (float, int)):
        if value < 0 or int(value) != value:
            raise InputError('El código numérico debe ser entero y no negativo.')
        if len(str(int(value))) > 15:
            raise InputError('Código numérico con más de 15 dígitos: recupere el original como texto.')
        mask = cell.number_format
        if not re.fullmatch(r'0+', mask or '') or len(mask) != length:
            raise InputError('Código numérico sin máscara compatible. Escríbalo como texto con sus ceros.')
        result = str(int(value)).zfill(length)
    elif isinstance(value, str):
        result = value.strip()
    else:
        raise InputError('El código debe ser texto de dígitos.')
    if not re.fullmatch(r'[0-9]+', result) or len(result) != length:
        raise InputError(f'El código debe tener {length} dígitos; no se truncó ni completó texto.')
    return result

def validate(content, rules=Rules()):
    wb = open_excel(content)
    try:
        if rules.sheet not in wb.sheetnames:
            raise InputError('La hoja indicada no existe. Indique su nombre exacto.')
        ws = wb[rules.sheet]
        if ws.max_row > 5001 or ws.max_column > 100:
            raise InputError('Máximo 5.000 filas y 100 columnas; quite formatos sobrantes o divida el libro.')
        if ws.merged_cells.ranges:
            raise InputError('La hoja contiene celdas combinadas. Use una tabla sin combinaciones.')
        mapping = {}
        seen = set()
        for cell in ws[1]:
            if cell.value is None:
                if any(ws.cell(r, cell.column).value is not None for r in range(2, ws.max_row + 1)):
                    raise InputError('Hay datos en una columna sin encabezado.')
                continue
            header = str(cell.value).strip()
            if header in seen:
                raise InputError('Hay encabezados duplicados.')
            seen.add(header)
            canonical = ALIASES.get(header)
            if canonical is None:
                raise InputError('Encabezado no reconocido. Consulte los alias de la guía; no se infirió un mapeo.')
            if canonical in mapping:
                raise InputError('Dos encabezados representan el mismo campo; el mapeo es ambiguo.')
            mapping[canonical] = cell.column
        if REQUIRED - mapping.keys():
            raise InputError('Faltan columnas: ' + ', '.join(sorted(REQUIRED - mapping.keys())))
        rows, blank = [], 0
        for index in range(2, ws.max_row + 1):
            if all(ws.cell(index, c).value is None or (isinstance(ws.cell(index,c).value,str) and not ws.cell(index,c).value.strip()) for c in range(1,ws.max_column+1)):
                blank += 1
                continue
            row = Row(index)
            def cell(key):
                return ws.cell(index, mapping[key]) if key in mapping else None
            try:
                row.code = code_value(cell('COD_ALUMNO'), rules.code_length)
                if isinstance(cell('COD_ALUMNO').value, (int, float)):
                    row.notes.append('Código reconstruido desde la máscara numérica del Excel.')
            except InputError as exc:
                row.issues.append(str(exc))
            for key, attr in [('NOMBRE_ESTUDIANTE','name'),('TITULO','title'),('DETALLE','detail'),('PRIORIDAD','priority')]:
                c = cell(key)
                value = c.value if c is not None else None
                if c is not None and c.data_type == 'f':
                    row.issues.append(f'{key}: no se admiten fórmulas.')
                elif value is not None and not isinstance(value, str):
                    row.issues.append(f'{key}: debe ser texto.')
                elif isinstance(value,str) and value.strip():
                    setattr(row, attr, value)
                elif key in REQUIRED:
                    row.issues.append(f'{key}: contenido obligatorio vacío.')
            if row.priority != 'Media':
                row.issues.append('Prioridad no reconocida. Solo Media está habilitada para diagnóstico.')
            if not row.name.strip():
                row.issues.append('Falta nombre para revisar destinatario; la regla de facultad sigue pendiente.')
            for key in ['COD_TRAMITE','GPA','CREDITOS_APROBADOS','ESTADO','MENSAJE_LOG']:
                c = cell(key)
                if c is not None and c.value not in (None, ''):
                    row.issues.append('Hay resultados previos. Concilie ese archivo antes de preparar una nueva entrada.')
                    break
            payload = [row.code,row.name,row.title,row.detail,row.priority]
            row.fingerprint = digest(json.dumps(payload, ensure_ascii=False).encode())
            rows.append(row)
        if not rows:
            raise InputError('El libro no contiene filas de trámites.')
        by_code, by_content = {}, {}
        for row in rows:
            if row.code:
                by_code.setdefault(row.code, []).append(row)
            by_content.setdefault(row.fingerprint, []).append(row)
        for group in by_code.values():
            if len({r.name.strip().casefold() for r in group}) > 1:
                for row in group:
                    row.issues.append('Un mismo código tiene nombres distintos. Confirme alumno o facultad; no se eliminó la fila.')
        for group in by_content.values():
            if len(group) > 1:
                for row in group:
                    row.issues.append('Posible duplicado de contenido completo; requiere revisión, sin eliminar filas.')
        return {'hash': digest(content), 'rules': asdict(rules), 'rows': [asdict(r) for r in rows], 'blank_rows': blank}
    finally:
        wb.close()

def export_results(content, rules, rows, run_id):
    wb = open_excel(content)
    ws = wb[rules.sheet]
    columns = ['ESTADO','MENSAJE_LOG','MODO_CAFI','RUN_ID_CAFI','REFERENCIA_SIMULADA']
    start = ws.max_column + 1
    for i, header in enumerate(columns, start):
        ws.cell(1,i,header).font = Font(bold=True,color='FFFFFF')
        ws.cell(1,i).fill = PatternFill('solid',fgColor='8A1538')
        ws.column_dimensions[ws.cell(1,i).column_letter].width = 27 if i != start+1 else 65
    for row in rows:
        values = [row['state'],row['message'],'SIMULACIÓN / SIN CONEXIÓN CAFI',run_id,row.get('transaction_id') or '']
        for i,value in enumerate(values,start):
            ws.cell(row['row'],i,value).alignment = Alignment(wrap_text=True,vertical='top')
    out = io.BytesIO()
    wb.save(out)
    wb.close()
    return out.getvalue()

def sample_book(incidents=False):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Hoja 1'
    ws.append(['COD. ALUMNO','NOMBRE ESTUDIANTE','TITULO','DETALLE','CÓDIGO:','GPA:','CRED.APROB:'])
    for i, title in enumerate(['INCOMPLETO','HOMOLOGACIÓN','EXAMEN SUPLETORIO'],1):
        ws.append([f'990000000{i}',f'Estudiante de prueba {i}',title,f'DATO SINTÉTICO. Solicitud de prueba {i}.\nSegunda línea con acentos: revisión académica.',None,None,None])
        ws.cell(i+1,1).number_format = '@'
    if incidents:
        ws['A3'] = ws['A2'].value
        ws['D4'] = None
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill('solid',fgColor='8A1538')
        cell.font = Font(color='FFFFFF',bold=True)
    for col,width in [('A',22),('B',30),('C',28),('D',65),('E',18),('F',12),('G',18)]:
        ws.column_dimensions[col].width=width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True,vertical='top')
        ws.row_dimensions[row[0].row].height=48
    output = io.BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()
