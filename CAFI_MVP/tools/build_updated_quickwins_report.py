from copy import deepcopy
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "Reporte de Avance Quickwins - Maria Emilia Final.docx"
OUTPUT = ROOT / "Reporte de Avance Quickwins - Maria Emilia Aguirre Beltrán Actualizado.docx"

NAVY = "1F4E78"
PALE_BLUE = "DDEBF7"
VERY_PALE_BLUE = "F3F7FB"
GRID = "D9D9D9"
BLACK = RGBColor(0, 0, 0)


def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=GRID, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=65, start=110, bottom=65, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def set_no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    tr_pr.append(node)


def set_column_widths(table, widths):
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width


def write_cell(cell, text, bold=False, color=BLACK, size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.color.rgb = color
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    set_cell_border(cell)


def add_table(doc, headers, rows, widths, font_size=9.3, first_col_bold=False):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    set_column_widths(table, widths)
    header = table.rows[0]
    set_repeat_table_header(header)
    for cell, text in zip(header.cells, headers):
        set_cell_shading(cell, NAVY)
        write_cell(cell, text, bold=True, color=RGBColor(255, 255, 255), size=font_size, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        set_no_split(table.rows[-1])
        for col_index, (cell, value) in enumerate(zip(cells, values)):
            if row_index % 2 == 1:
                set_cell_shading(cell, VERY_PALE_BLUE)
            write_cell(cell, value, bold=first_col_bold and col_index == 0, size=font_size,
                       align=WD_ALIGN_PARAGRAPH.CENTER if len(value) < 18 and col_index in (0, 3, 4) else WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_title(doc, title, subtitle):
    p = doc.add_paragraph(style="Title")
    p_pr = p._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title)
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(22)
    run.bold = True
    run.font.color.rgb = BLACK
    p2 = doc.add_paragraph(style="Subtitle")
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(16)
    r2 = p2.add_run(subtitle)
    r2.font.name = "Arial"
    r2._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    r2._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    r2.font.size = Pt(13)
    r2.font.color.rgb = BLACK
    r2.italic = False


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.color.rgb = BLACK
    return p


def add_labeled(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(label + " ")
    r.bold = True
    p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.05
        p.add_run(item)


def add_formula(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    r.bold = True
    return p


def build():
    copy2(TEMPLATE, OUTPUT)
    doc = Document(OUTPUT)
    clear_body(doc)

    section = doc.sections[0]
    section.top_margin = Inches(0.60)
    section.bottom_margin = Inches(0.60)
    section.left_margin = Inches(0.70)
    section.right_margin = Inches(0.70)
    footer = section.footer
    footer.paragraphs[0].text = "Levantamiento, seguimiento e impacto de soluciones de automatización"
    footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.paragraphs[0].runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(100, 100, 100)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.05
    title_p_pr = doc.styles["Title"]._element.find(qn("w:pPr"))
    if title_p_pr is not None:
        title_border = title_p_pr.find(qn("w:pBdr"))
        if title_border is not None:
            title_p_pr.remove(title_border)
    for name, size in (("Heading 1", 15), ("Heading 2", 12.5)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK

    add_title(doc, "REPORTE DE AVANCE QUICKWINS", "Automatización del ingreso de trámites en SIAC CAFI")
    add_labeled(doc, "Objetivo:", "Documentar el diagnóstico, el avance técnico, el plan de validación y la medición de impacto de la automatización del ingreso de trámites desde una agenda estructurada hacia SIAC, módulo CAFI.")
    add_table(doc, ["Campo", "Información"], [
        ("Nombre del colaborador", "María Emilia Aguirre Beltrán"),
        ("Área", "Facultad de Comunicación, Universidad Espíritu Santo UEES"),
        ("Estudiante responsable", "Melany Brito"),
        ("Fecha de corte", "22 de septiembre de 2026"),
    ], [Inches(2.55), Inches(4.45)], first_col_bold=True)

    add_heading(doc, "1. Identificación del punto de dolor")
    add_labeled(doc, "Proceso a mejorar:", "ingreso, registro, impresión y actualización de resultados de trámites estudiantiles en SIAC, módulo CAFI.")
    add_labeled(doc, "Problema principal:", "Saskya Torres Guerrero debe trasladar manualmente los datos de cada fila de la agenda a CAFI, esperar la respuesta del sistema, imprimir el trámite y volver a la agenda para registrar el código de trámite, el GPA y los créditos aprobados.")
    add_labeled(doc, "Frecuencia:", "cada dos semanas. Para el indicador mensual de Quick Wins se utiliza una referencia de dos ciclos por mes; la frecuencia anual debe confirmarse con el calendario operativo de la Facultad.")
    add_labeled(doc, "Tiempo actual:", "12 horas por ciclo para aproximadamente 45 registros, realizadas de forma intermitente. La referencia equivale a 16 minutos manuales por registro y a 24 horas mensuales por facultad bajo dos ciclos mensuales.")
    add_labeled(doc, "Personas involucradas:", "en la Facultad de Comunicación, Saskya Torres Guerrero ejecuta el proceso y María Emilia Aguirre Beltrán valida los resultados consolidados del lote. El escenario de escalabilidad considera una persona administrativa operativa por cada una de las ocho facultades UEES.")
    add_labeled(doc, "Puntos de dolor identificados:", "")
    add_bullets(doc, [
        "Transcripción repetitiva de código, título, detalle y variables de cada trámite, con riesgo de digitación, omisión o asociación incorrecta.",
        "Navegación recurrente por SIAC, CAFI, Procesos y Trámites Estudiante para cada fila.",
        "Tiempos de respuesta variables del sistema y necesidad de comprobar visualmente el mensaje de éxito antes de continuar.",
        "Impresión individual y lectura manual del código de trámite, GPA y créditos aprobados.",
        "Actualización posterior de la agenda, con riesgo de pérdida de trazabilidad o duplicación si un guardado queda incierto.",
    ])

    add_heading(doc, "2. Datos de entrada")
    add_labeled(doc, "Información necesaria:", "código de alumno o facultad, nombre del estudiante para la validación, título, detalle y los campos que estén definidos en cada fila de la agenda.")
    add_labeled(doc, "Origen:", "agenda principal gestionada en Google Sheets. En el MVP local se trabajará inicialmente con una exportación Excel de la misma estructura.")
    add_labeled(doc, "Formato:", "Google Sheets o Excel. La hoja de prueba mantiene los encabezados COD. ALUMNO, NOMBRE ESTUDIANTE, TITULO, DETALLE, CÓDIGO, GPA y CRED.APROB.")
    add_labeled(doc, "Variabilidad:", "los títulos y detalles son variables por fila; el RPA debe transferirlos de forma literal y no imponer un catálogo fijo.")
    add_labeled(doc, "Preparación y validación:", "se deben conservar los ceros iniciales de los códigos, revisar campos obligatorios, detener filas incompletas o inconsistentes y bloquear cualquier discrepancia entre el destinatario visible en CAFI y la fila de origen.")

    add_heading(doc, "3. Herramientas y plataformas que intervienen")
    add_table(doc, ["Herramienta", "Uso actual", "Uso previsto en la solución"], [
        ("Google Sheets / Excel", "Agenda y fuente de datos.", "Entrada validada y destino de los resultados por fila en la fase correspondiente."),
        ("SIAC - CAFI", "Registro, confirmación e impresión de trámites.", "Aplicación de escritorio operada dentro de una sesión Windows autorizada."),
        ("RPA local", "No existe en el proceso manual.", "Ventana nativa, validación Excel, ejecución secuencial, bitácora y control de estados."),
        ("Python, pywinauto, UIA y Win32", "No participan en la ejecución manual.", "Lectura e interacción con controles Windows observados en la laptop de destino."),
    ], [Inches(1.45), Inches(2.35), Inches(3.2)])
    add_labeled(doc, "Orden de uso:", "agenda validada → SIAC/CAFI → confirmación de guardado → impresión individual → cierre del reporte → siguiente fila. Al terminar el lote, María Emilia revisa los resultados consolidados y se actualiza la copia de resultados.")
    add_labeled(doc, "Accesos y limitaciones:", "SIAC/CAFI requiere una sesión institucional autorizada en la laptop Windows de Saskya. El RPA y SIAC deben ejecutarse en la misma sesión de Windows y con el mismo nivel de permisos. No se ha verificado una API institucional para este flujo; por ello, los controles deben inspeccionarse directamente en el equipo donde opera SIAC.")

    add_heading(doc, "4. Proceso actual paso a paso")
    add_table(doc, ["Paso", "Actividad", "Herramienta", "Responsable", "Tiempo", "Observaciones y dolor"], [
        ("1", "Revisar la agenda y validar los datos de la fila.", "Sheets/Excel", "Saskya", "Variable", "Se deben identificar datos faltantes, duplicados o inconsistentes antes de iniciar."),
        ("2", "Abrir SIAC y acceder a CAFI.", "SIAC", "Saskya", "Variable", "Depende de la sesión institucional y de la disponibilidad del sistema."),
        ("3", "Abrir Trámites Estudiante y crear un nuevo proceso.", "CAFI", "Saskya", "Variable", "La navegación se repite por cada trámite."),
        ("4", "Ingresar código, título y detalle; comprobar el destinatario resuelto.", "CAFI", "Saskya", "Variable", "Existe riesgo de asociar el trámite a un destinatario incorrecto."),
        ("5", "Guardar una sola vez y esperar el mensaje de éxito.", "CAFI", "Saskya", "Variable", "No debe reenviarse un trámite cuando el resultado sea incierto."),
        ("6", "Imprimir, verificar el reporte y recuperar código de trámite, GPA y créditos.", "CAFI", "Saskya", "Variable", "La impresión y la lectura se realizan individualmente."),
        ("7", "Actualizar la agenda y dejar el formulario listo para la siguiente fila.", "Sheets/Excel", "Saskya", "Variable", "La transferencia manual adicional puede producir inconsistencias."),
    ], [Inches(.48), Inches(1.2), Inches(.88), Inches(.85), Inches(.68), Inches(2.91)], font_size=8.4)

    add_heading(doc, "5. Formato de salida y uso de la información")
    add_labeled(doc, "Resultado final actual:", "trámite registrado e impreso en CAFI, con código de trámite, GPA y créditos aprobados disponibles para su registro en la agenda.")
    add_labeled(doc, "Salida del MVP:", "una copia Excel de resultados, denominada resultados_cafi.xlsx, que conserva el archivo fuente y registra código de trámite, GPA, créditos, estado y observación al concluir el lote y después de la revisión funcional consolidada.")
    add_labeled(doc, "Usuario y receptor principal:", "Saskya Torres Guerrero, asistente administrativa de la Facultad de Comunicación.")
    add_labeled(doc, "Uso posterior:", "los tres resultados deben alimentar las columnas correspondientes de la agenda. La actualización directa de Google Sheets queda como fase posterior, después de validar el ciclo completo con Excel local y definir cuenta, permisos y método de autenticación.")
    add_labeled(doc, "Criterios de corrección:", "cada fila se registra una sola vez; el destinatario visible coincide con la fila; se comprueba el éxito del guardado; la impresión corresponde al trámite; y los tres resultados se registran sin alterar los datos de origen.")

    add_heading(doc, "6. Responsables e involucrados")
    add_table(doc, ["Rol", "Responsabilidad"], [
        ("Saskya Torres Guerrero", "Responsable operativa y usuaria principal en la Facultad de Comunicación. Es la persona autorizada para crear los trámites en SIAC/CAFI y participará en el piloto."),
        ("María Emilia Aguirre Beltrán", "Responsable de Quick Wins y validadora funcional. Revisa los resultados consolidados al finalizar el lote antes de actualizar el Excel."),
        ("Melany Brito", "Apoyo y consultoría principal para levantamiento, organización documental, definición de prompts y desarrollo técnico en Codex."),
        ("Facultad de Comunicación", "Área de aplicación inicial y validación del proceso."),
    ], [Inches(2.2), Inches(4.8)])

    add_heading(doc, "7. Alcance de la solución")
    add_labeled(doc, "Punto específico que se ataca:", "la transcripción y registro repetitivo de la agenda hacia SIAC/CAFI, incluida la impresión individual y la devolución trazable de los tres resultados.")
    add_labeled(doc, "Inicio y final del flujo automatizado:", "el flujo inicia con un Excel validado y SIAC abierto en la sesión autorizada. Para cada fila, registra, imprime, cierra el reporte y continúa con la siguiente. Finaliza cuando se completa el lote, se revisan los resultados consolidados y se actualiza la copia Excel.")
    add_labeled(doc, "Fuera de alcance actual:", "la lectura de correos, la aprobación de solicitudes académicas, la modificación de títulos o detalles, la conexión directa a Google Sheets y cualquier reintento automático ante un guardado incierto.")

    add_heading(doc, "8. Tipo de solución seleccionada")
    add_labeled(doc, "Solución seleccionada:", "RPA atendido y local en Windows, empaquetado como aplicación de escritorio con ventana nativa e iniciador EXE o BAT. No utiliza una página HTML ni requiere abrir Visual Studio Code para la operación.")
    add_labeled(doc, "Conexiones necesarias:", "Excel local, SIAC/CAFI, controles Windows accesibles mediante UIA o Win32 y el diálogo de impresión o reporte. Google Sheets se conectará solo después de una validación funcional y de disponer de permisos autorizados.")
    add_labeled(doc, "Controles de seguridad y continuidad:", "conservación del Excel original, bitácora por fila, persistencia de intención antes de guardar, separación entre guardado e impresión, detención ante resultados inciertos y revisión funcional consolidada antes de exportar los tres resultados.")
    add_labeled(doc, "Acciones que conserva la persona:", "validar la agenda, abrir o autorizar la sesión de SIAC, cargar el Excel, autorizar filas de piloto, revisar los resultados consolidados al finalizar el lote y atender cualquier excepción. La laptop queda reservada para el RPA durante la ejecución y no debe utilizarse en paralelo para otras tareas.")

    add_heading(doc, "9. Avance del proyecto")
    add_labeled(doc, "Reuniones realizadas:", "cinco reuniones de aproximadamente una hora. El 8 de septiembre de 2026 se realizó el diagnóstico y levantamiento; el 11 de septiembre de 2026 se organizaron las fuentes, se definieron prompts y se inició el desarrollo del MVP en Codex. Las reuniones posteriores profundizaron la prueba en Windows, el flujo de impresión, la confirmación humana y la medición de impacto.")
    add_labeled(doc, "Desarrollo completado al corte:", "se estructuró la versión 0.4.0 del MVP como aplicación nativa de Windows con iniciadores EXE y BAT; validación de Excel; preservación de ceros iniciales; bitácora y control de estados; espera configurable entre 10 y 15 segundos, con 12 segundos como valor inicial; continuidad automática entre filas tras cada impresión; y exportación a una copia de resultados al cierre del lote.")
    add_labeled(doc, "Avances del último piloto:", "se ejecutó INICIAR_CAFI.bat, se inspeccionó SIAC mediante UIA y Win32, se identificaron las ventanas Procesos e Ingreso de Procesos y se comprobó manualmente una fila sintética. El trámite 105000 se guardó una sola vez y el reporte mostró GPA 0 y créditos aprobados 0; estos datos fueron confirmados para el Excel de resultados. El piloto fue manual y no demuestra todavía un ciclo completo ejecutado autónomamente por el RPA.")
    add_labeled(doc, "Pruebas de desarrollo:", "la documentación del último avance registra 58 pruebas locales aprobadas. Estas pruebas cubren el núcleo, contratos del controlador y casos simulados; no acreditan todavía un ciclo completo de escritura, guardado, impresión y extracción ejecutado autónomamente dentro de SIAC real.")
    add_labeled(doc, "Preparación del piloto:", "se elaboró un prompt operativo para inspeccionar los controles de impresión y probar una fila autorizada antes de ejecutar dos o tres filas consecutivas. También se creó un Excel sintético de tres filas para validar la estructura de entrada sin enviar datos ficticios a SIAC.")
    add_labeled(doc, "Pendientes críticos:", "completar el perfil de controles en la laptop Windows de Saskya. Permanecen 16 de 18 controles obligatorios agrupados en formulario, guardado, impresión, resultados y preparación de la siguiente fila; después se debe completar un ciclo autónomo de una fila y comprobar dos o tres filas consecutivas sin duplicación.")
    add_labeled(doc, "Pruebas con información real:", "no existe evidencia de trámites reales guardados por el RPA. La carga productiva continúa pendiente de un piloto supervisado y autorizado.")

    add_heading(doc, "10. Medición de impacto ahorro de horas")
    add_heading(doc, "10.1 Variables disponibles y pendientes", level=2)
    add_table(doc, ["Variable", "Estado al 22 de septiembre de 2026"], [
        ("Tiempo actual por ejecución", "12 horas por lote de aproximadamente 45 registros, equivalentes a 720 minutos o 16 minutos por registro."),
        ("Tiempo técnico estimado con RPA", "Primer registro: 2 min 12 s. Cada registro siguiente: 1 min 24 s. Para 45 registros: 1 h 03 min 48 s, sin intervención entre filas."),
        ("Intervención humana estimada", "3 minutos iniciales para sesión y carga del Excel. El RPA no solicita confirmación por trámite: imprime, cierra el reporte y continúa automáticamente con la siguiente fila. La revisión funcional se realiza al finalizar el lote."),
        ("Tiempo total estimado del RPA", "1 h 06 min 48 s por lote de 45 registros, al sumar los 3 minutos iniciales al tiempo técnico. Debe validarse en piloto."),
        ("Frecuencia", "2 ciclos mensuales de referencia y 8 ciclos por semestre; para el año académico se usan 16 ciclos."),
        ("Personas involucradas", "Escenario de escala: 8 personas operativas, una por cada facultad UEES."),
        ("Costo por hora laboral", "$5, variable fija establecida por la plantilla Quick Wins."),
    ], [Inches(2.35), Inches(4.65)], first_col_bold=True)
    add_heading(doc, "10.2 a 10.6 Fórmulas aplicables", level=2)
    add_formula(doc, "Escenario Saskya - ahorro por lote = 12 h - 3 min = 11 h 57 min, equivalentes a 11,95 horas por ciclo en la Facultad de Comunicación.")
    add_formula(doc, "Escenario Saskya - ahorro mensual = (11,95 h x 2 ciclos) = 23,9 horas; ahorro económico mensual = 23,9 x $5 = $119,50.")
    add_formula(doc, "Escenario Saskya - ahorro anual = 11,95 h x 16 ciclos = 191,2 horas; ahorro económico anual = 191,2 x $5 = $956.")
    add_formula(doc, "Escenario escalado a 8 facultades - horas actuales por mes = (720 min x 2 ciclos x 8 personas) / 60 = 192 horas.")
    add_formula(doc, "Escenario escalado a 8 facultades - ahorro mensual = 23,9 h x 8 = 191,2 horas; ahorro económico mensual = $119,50 x 8 = $956.")
    add_formula(doc, "Escenario escalado a 8 facultades - ahorro anual = 191,2 h x 8 = 1.529,6 horas; ahorro económico anual = $956 x 8 = $7.648.")

    add_heading(doc, "11. Conversión del ahorro de tiempo a dinero")
    add_labeled(doc, "Fórmula:", "Ahorro económico = Horas humanas ahorradas x $5 por hora. Esta valorización utiliza el valor fijo de la plantilla Quick Wins.")
    add_labeled(doc, "Línea base y proyección:", "para Saskya, el ahorro estimado es de 23,9 horas mensuales o $119,50, y de 191,2 horas anuales o $956. Si se escala a las ocho facultades, el ahorro estimado asciende a 191,2 horas mensuales o $956, y a 1.529,6 horas anuales o $7.648. Son proyecciones preliminares y deben confirmarse con el piloto.")

    add_heading(doc, "12. Cálculo del ROI")
    add_labeled(doc, "Fórmula Quick Wins:", "ROI (%) = [(Horas ahorradas acumuladas - Horas totales de desarrollo) / Horas totales de desarrollo] x 100.")
    add_labeled(doc, "Aplicación al corte:", "las 12 horas reportadas corresponden al tiempo manual de operación por ciclo, por lo que no deben utilizarse como inversión de desarrollo en el denominador. Se registran cinco horas documentadas de reuniones. Para completar la medición mientras se consolida la bitácora, se estima una inversión total de 25 horas: 5 horas de reuniones y 20 horas adicionales de levantamiento, documentación, desarrollo, pruebas, correcciones y preparación del piloto. Es una estimación de gestión, no una medición definitiva.")
    add_table(doc, ["Variable de ROI", "Estado y dato requerido"], [
        ("Horas ahorradas acumuladas", "Saskya: 11,95 horas por ciclo y 191,2 horas en 16 ciclos. Escala a 8 facultades: 95,6 horas por ciclo y 1.529,6 horas en 16 ciclos. El acumulado real debe contabilizar únicamente ciclos ejecutados y validados."),
        ("Horas totales de desarrollo", "25 horas estimadas para el escenario base: 5 horas de reuniones y 20 horas de trabajo técnico y de implementación. Debe sustituirse por la bitácora real cuando finalice el piloto."),
        ("ROI Saskya por ciclo", "[(11,95 - 25) / 25] x 100 = -52,2 %. La inversión todavía no se recupera en un solo ciclo individual."),
        ("ROI Saskya anual", "[(191,2 - 25) / 25] x 100 = 664,8 % en 16 ciclos académicos."),
        ("ROI escalado por ciclo", "[(95,6 - 25) / 25] x 100 = 282,4 % cuando participan las ocho facultades."),
        ("ROI escalado anual", "[(1.529,6 - 25) / 25] x 100 = 6.018,4 % en 16 ciclos académicos."),
        ("ROI económico", "Saskya: ahorro anual $956; ROI = [($956 - $125) / $125] x 100 = 664,8 %. Escala a 8 facultades: ahorro anual $7.648; ROI = [($7.648 - $125) / $125] x 100 = 6.018,4 %. Debe sustituirse por el costo real."),
    ], [Inches(2.35), Inches(4.65)], first_col_bold=True)

    add_heading(doc, "13. Resumen final de impacto")
    add_table(doc, ["Indicador", "Resultado al 22 de septiembre de 2026"], [
        ("Tiempo actual por ejecución", "12 horas por lote de aproximadamente 45 registros."),
        ("Tiempo técnico estimado con RPA", "1 h 03 min 48 s por 45 registros, sin intervención entre filas."),
        ("Intervención humana estimada", "3 minutos por lote para abrir o autorizar SIAC y cargar el Excel. La revisión funcional ocurre al finalizar el lote."),
        ("Tiempo total estimado del RPA", "1 h 06 min 48 s por lote de 45 registros, sujeto a validación del piloto."),
        ("Ejecuciones", "2 ciclos mensuales de referencia; 16 ciclos por año académico."),
        ("Personas involucradas", "Saskya es la usuaria inicial; el escenario de escala considera 8 personas, una por facultad UEES."),
        ("Ahorro mensual de Saskya", "23,9 horas y $119,50 estimados para 2 ciclos en la Facultad de Comunicación."),
        ("Ahorro anual de Saskya", "191,2 horas y $956 estimados para 16 ciclos académicos."),
        ("Ahorro mensual escalado", "191,2 horas y $956 estimados para las 8 facultades."),
        ("Ahorro anual escalado", "1.529,6 horas y $7.648 estimados para 16 ciclos académicos."),
        ("Inversión estimada de desarrollo", "25 horas en el escenario base: 5 horas documentadas de reuniones y 20 horas estimadas de trabajo técnico."),
        ("ROI estimado de Saskya por ciclo", "-52,2 %, porque la inversión estimada de 25 horas aún no se recupera en un ciclo individual."),
        ("ROI estimado de Saskya anual", "664,8 %, usando 191,2 horas ahorradas y 25 horas de inversión."),
        ("ROI estimado escalado por ciclo", "282,4 %, usando 95,6 horas ahorradas por ciclo en las ocho facultades."),
        ("ROI estimado escalado anual", "6.018,4 %, usando 1.529,6 horas ahorradas y 25 horas de inversión."),
    ], [Inches(2.35), Inches(4.65)], first_col_bold=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
