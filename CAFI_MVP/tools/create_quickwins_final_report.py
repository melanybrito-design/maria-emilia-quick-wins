from copy import deepcopy
from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path('/Users/melanybrito/Desktop/Maria Emilia (QuickWins) /CAFI_MVP')
TEMPLATE = Path('/Users/melanybrito/Desktop/Reporte de Avance Quickwins.docx')
OUTPUT = ROOT / 'Reporte de Avance Quickwins - Maria Emilia Final.docx'
BLUE = '1F4E79'
LIGHT_BLUE = 'D9EAF7'


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = 'Times New Roman'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    r.font.size = Pt(10.5)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table(table, headers, widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for i, heading in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_text(cell, heading, bold=True, color='FFFFFF')
        shade(cell, BLUE)
        if widths:
            cell.width = Inches(widths[i])
    for row in table.rows[1:]:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = 'Heading 1' if level == 1 else 'Heading 2'
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = 'Arial'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    r.font.color.rgb = RGBColor.from_string(BLUE)
    r.bold = True
    return p


def add_text(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.08
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def add_answer(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(label + ' ')
    r.bold = True
    p.add_run(value)


def add_key_value_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, (key, value) in enumerate(rows):
        set_cell_text(table.rows[idx].cells[0], key, bold=True)
        set_cell_text(table.rows[idx].cells[1], value)
        shade(table.rows[idx].cells[0], LIGHT_BLUE)
    return table


def main():
    copy2(TEMPLATE, OUTPUT)
    doc = Document(OUTPUT)
    body = doc._element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'):
            body.remove(child)

    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    normal.font.size = Pt(11)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run('REPORTE DE AVANCE QUICKWINS')
    r.bold = True
    r.font.name = 'Arial'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    r.font.size = Pt(25)
    r.font.color.rgb = RGBColor.from_string('203E67')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Levantamiento y avance de automatización del proceso CAFI')
    r.font.name = 'Arial'
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor.from_string('2F6193')
    add_text(doc, 'Objetivo: Documentar el diagnóstico, el avance técnico y la medición de impacto de la automatización del ingreso de trámites desde Google Sheets o Excel hacia SIAC–CAFI.', 'Objetivo:')
    add_key_value_table(doc, [
        ('Nombre del colaborador:', 'Saskya Torres Guerrero'),
        ('Área:', 'Facultad de Comunicación, Universidad Espíritu Santo UEES'),
        ('Estudiante/s responsable/s:', 'María Emilia Aguirre Beltrán (Quick Wins) y Melany Brito (apoyo y consultoría principal).'),
        ('Fecha de corte:', '11 de septiembre de 2026'),
    ])

    add_heading(doc, '1. Identificación del punto de dolor')
    add_answer(doc, 'Proceso a mejorar:', 'ingreso, registro, impresión y actualización de resultados de trámites estudiantiles en SIAC, módulo CAFI.')
    add_answer(doc, 'Problema principal:', 'la asistente debe trasladar manualmente los datos de cada fila desde la agenda hacia CAFI, esperar la respuesta del sistema, imprimir cada trámite y regresar a la agenda para registrar el código de trámite, GPA y créditos aprobados.')
    add_answer(doc, 'Frecuencia:', 'cada dos semanas; para la medición se usa una referencia de dos ciclos mensuales.')
    add_answer(doc, 'Tiempo actual:', '12 horas por ciclo, ejecutadas de manera intermitente. La línea base mensual es de 24 horas.')
    add_answer(doc, 'Personas involucradas:', 'una persona operativa principal: Saskya Torres Guerrero.')
    add_text(doc, 'Puntos de dolor identificados:', 'Puntos de dolor identificados:')
    add_bullets(doc, [
        'Transcripción repetitiva de código del estudiante, título, detalle y demás variables, con riesgo de digitación u omisión.',
        'Navegación recurrente por SIAC, CAFI, Procesos y Trámites Estudiante para cada registro.',
        'Tiempos de espera variables y necesidad de comprobar visualmente que el sistema confirmó el guardado.',
        'Impresión individual y lectura manual de código de trámite, GPA y créditos aprobados.',
        'Actualización posterior de la agenda principal, con riesgo de perder trazabilidad o duplicar una carga ante una respuesta incierta.',
    ])

    add_heading(doc, '2. Datos de entrada')
    add_answer(doc, 'Información necesaria:', 'datos del estudiante, código del estudiante, título, detalle y las variables que estén definidas en cada fila de la agenda.')
    add_answer(doc, 'Origen:', 'agenda principal administrada en Google Sheets; para el RPA local se podrá usar su exportación o un archivo Excel con la misma estructura.')
    add_answer(doc, 'Formato:', 'Google Sheets y Excel.')
    add_answer(doc, 'Estructura:', 'los títulos y datos son variables y dependen de cada fila; el RPA no debe imponer un catálogo fijo ni modificar los valores fuente.')
    add_answer(doc, 'Validaciones requeridas:', 'conservar ceros iniciales y contenido literal; verificar campos obligatorios, filas incompletas, duplicadas o inconsistentes antes de iniciar el lote.')

    add_heading(doc, '3. Herramientas y plataformas que intervienen')
    tbl = doc.add_table(rows=1, cols=3)
    style_table(tbl, ['Herramienta', 'Uso actual', 'Uso previsto en la solución'], [1.6, 2.5, 2.7])
    for values in [
        ('Google Sheets / Excel', 'Agenda y fuente de datos.', 'Fuente validada y destino de los resultados por fila.'),
        ('SIAC – CAFI', 'Registro, confirmación e impresión de trámites.', 'Aplicación de escritorio operada dentro de la sesión Windows autorizada.'),
        ('RPA local', 'No existe en el proceso manual.', 'Ventana nativa, lectura de archivo, control de estados y bitácora.'),
        ('Python, pywinauto, UIA y Win32', 'No participan en la operación manual.', 'Automatización de controles Windows y lectura verificable de la interfaz.'),
    ]:
        row = tbl.add_row().cells
        for c, v in zip(row, values): set_cell_text(c, v)
    add_answer(doc, 'Orden de uso:', 'agenda → SIAC/CAFI → confirmación e impresión → extracción de resultados → actualización de la agenda.')
    add_answer(doc, 'Accesos y limitaciones:', 'SIAC/CAFI requiere sesión y permisos autorizados en la laptop Windows de Saskya. No existe evidencia de una API institucional para este flujo; la automatización se implementa sobre la interfaz de escritorio y debe validar sus controles reales.')

    add_heading(doc, '4. Proceso actual: paso a paso')
    tbl = doc.add_table(rows=1, cols=6)
    style_table(tbl, ['Paso', 'Actividad', 'Herramienta', 'Responsable', 'Tiempo', 'Observaciones / dolor'], [0.45, 1.55, 1.0, 1.0, 0.75, 1.7])
    steps = [
        ('1', 'Revisar la agenda y validar los datos de la fila.', 'Sheets/Excel', 'Saskya', 'Variable', 'Se deben detectar datos faltantes o inconsistentes.'),
        ('2', 'Ingresar a SIAC y abrir CAFI.', 'SIAC', 'Saskya', 'Variable', 'Depende de credenciales y disponibilidad de la sesión.'),
        ('3', 'Abrir Trámites Estudiante y crear un nuevo proceso.', 'CAFI', 'Saskya', 'Variable', 'Navegación repetitiva por cada trámite.'),
        ('4', 'Registrar código, título, detalle y prioridad según la agenda.', 'CAFI', 'Saskya', 'Variable', 'Riesgo de digitación o de asociar un estudiante incorrecto.'),
        ('5', 'Guardar y esperar la confirmación.', 'CAFI', 'Saskya', 'Variable', 'El sistema puede tardar; no se debe reenviar ante incertidumbre.'),
        ('6', 'Imprimir y recuperar código de trámite, GPA y créditos.', 'CAFI', 'Saskya', 'Variable', 'La impresión y lectura se hacen por trámite.'),
        ('7', 'Actualizar la agenda principal y continuar con la siguiente fila.', 'Sheets/Excel', 'Saskya', 'Variable', 'Transferencia manual adicional y riesgo de inconsistencia.'),
    ]
    for values in steps:
        row = tbl.add_row().cells
        for c, v in zip(row, values): set_cell_text(c, v)

    add_heading(doc, '5. Formato de salida y uso de la información')
    add_answer(doc, 'Resultado final actual:', 'trámite registrado e impreso en CAFI, con código de trámite, GPA y créditos aprobados disponibles para la actualización de la agenda.')
    add_answer(doc, 'Formato de salida:', 'registro en SIAC/CAFI, impresión individual y actualización de las columnas correspondientes de Google Sheets o Excel.')
    add_answer(doc, 'Usuario y receptor principal:', 'Saskya Torres Guerrero, asistente administrativa de la Facultad de Comunicación.')
    add_answer(doc, 'Criterios de corrección:', 'cada fila se registra una sola vez; el estudiante visible coincide con la fila; la impresión corresponde al trámite; y código, GPA y créditos se escriben en las columnas correctas sin alterar los datos fuente.')

    add_heading(doc, '6. Responsables e involucrados')
    tbl = doc.add_table(rows=1, cols=2)
    style_table(tbl, ['Rol', 'Responsabilidad'], [2.1, 5.9])
    for values in [
        ('Saskya Torres Guerrero', 'Responsable operativa y usuaria principal. Ejecuta actualmente el proceso manual y participará en la validación operativa.'),
        ('María Emilia Aguirre Beltrán', 'Responsable de Quick Wins. Impulsa y valida el proyecto de automatización.'),
        ('Melany Brito', 'Apoyo y consultoría principal para levantamiento, organización documental, definición de prompts y desarrollo técnico en Codex.'),
        ('Facultad de Comunicación', 'Área de aplicación inicial y validación del proceso.'),
    ]:
        row = tbl.add_row().cells
        for c, v in zip(row, values): set_cell_text(c, v)

    add_heading(doc, '7. Alcance de la solución')
    add_text(doc, 'La solución seleccionada es un RPA atendido y local. Inicia con la agenda validada y la sesión de SIAC disponible; toma los campos de cada fila, registra el trámite, espera y verifica la confirmación, ejecuta la impresión individual, recupera código de trámite, GPA y créditos aprobados, y prepara la actualización de la agenda principal.')
    add_text(doc, 'La actualización directa de Google Sheets es un requisito funcional. Su conexión se habilitará después de definir la cuenta, permisos, hoja y método autorizado. El RPA no aprueba solicitudes académicas, no decide contenidos, no cambia títulos ni sustituye la revisión humana ante inconsistencias.')

    add_heading(doc, '8. Tipo de solución seleccionada')
    add_answer(doc, 'Tipo de solución:', 'RPA local en Windows, empaquetable como aplicación de escritorio con ventana nativa para seleccionar el Excel e iniciar el proceso.')
    add_answer(doc, 'Conexiones necesarias:', 'archivo Excel o agenda exportada, SIAC–CAFI, impresora configurada y, en la fase de actualización, Google Sheets con acceso autorizado.')
    add_answer(doc, 'Acciones que conserva la persona:', 'validar la agenda, abrir o autorizar la sesión requerida, atender incidencias y validar el piloto antes de operar lotes.')
    add_bullets(doc, [
        'Mantener una bitácora por fila y separar estados de guardado, impresión, lectura y actualización.',
        'Detener el caso ante una confirmación incompleta; nunca reenviar automáticamente un trámite incierto.',
        'Verificar que el destinatario mostrado por CAFI corresponde a la fila antes de guardar.',
    ])

    add_heading(doc, '9. Avance del proyecto')
    add_answer(doc, 'Reuniones realizadas:', 'dos. El 8 de septiembre de 2026 se realizó el diagnóstico y levantamiento; el 11 de septiembre de 2026 se organizó la documentación, se estructuraron prompts y se inició el desarrollo en Codex y la preparación de Windows.')
    add_answer(doc, 'Funcionalidades desarrolladas:', 'base de RPA local con interfaz de escritorio, validación de archivos, bitácora, control de estados y lógica de recuperación.')
    add_answer(doc, 'Evidencia técnica:', 'conexión de lectura con SIAC mediante UIA y Win32 con estado CONECTADO_LECTURA, inventarios de controles generados y 40 pruebas ejecutadas en Windows.')
    add_answer(doc, 'Pendientes:', 'mapear los controles reales de Ingreso de Procesos, validar acceso y navegación, completar carga, impresión, lectura de resultados e integración autorizada con Google Sheets.')
    add_answer(doc, 'Pruebas con información real:', 'no se han guardado trámites con el RPA. La carga permanece deshabilitada hasta ejecutar un piloto supervisado y autorizado.')

    add_heading(doc, '10. Medición de impacto: ahorro de horas')
    add_heading(doc, '10.1 Variables necesarias', level=2)
    add_key_value_table(doc, [
        ('Tiempo actual por ejecución', '12 horas por ciclo quincenal (720 minutos).'),
        ('Tiempo con la solución por ejecución', 'Pendiente de medir en piloto con SIAC–CAFI.'),
        ('Ejecuciones mensuales', '2 ciclos de referencia.'),
        ('Personas involucradas', '1 persona operativa.'),
        ('Costo por hora laboral', '$5 (variable fija de la plantilla Quick Wins).'),
    ])
    add_heading(doc, '10.2 a 10.6 Fórmulas aplicables', level=2)
    add_text(doc, 'Tiempo ahorrado por ejecución = Tiempo actual − Tiempo con la solución.')
    add_text(doc, 'Horas actuales/mes = (Tiempo actual en min × Ejecuciones mensuales × Personas involucradas) ÷ 60 = (720 × 2 × 1) ÷ 60 = 24 horas.')
    add_text(doc, 'Horas con solución/mes = (Tiempo con solución en min × Ejecuciones mensuales × Personas involucradas) ÷ 60.')
    add_text(doc, 'Horas ahorradas/mes = Horas actuales/mes − Horas con solución/mes.')
    add_text(doc, 'Horas ahorradas/año = Horas ahorradas/mes × 12.')

    add_heading(doc, '11. Conversión del ahorro de tiempo a dinero')
    add_text(doc, 'Ahorro económico mensual = Horas ahorradas/mes × $5. Ahorro económico anual = Ahorro económico mensual × 12.')
    add_text(doc, 'La línea base de 24 horas mensuales equivale a $120 mensuales y $1.440 anuales de esfuerzo actual. No se presenta como ahorro: el ahorro se calculará cuando el piloto mida el tiempo con solución y el tiempo humano residual.')

    add_heading(doc, '12. Cálculo del ROI')
    add_text(doc, 'La fórmula requerida por Quick Wins se aplica en horas y compara el beneficio acumulado frente a las horas invertidas en desarrollar o implementar la solución.')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('ROI (%) = [(Horas ahorradas acumuladas − Horas totales de desarrollo) ÷ Horas totales de desarrollo] × 100')
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(BLUE)
    add_text(doc, 'Aplicación al corte: las 12 horas declaradas corresponden al tiempo manual del proceso y no a la inversión de desarrollo. Por tanto, no se sustituye ese dato en el denominador de ROI. El valor numérico de ROI queda pendiente hasta consolidar la bitácora de horas de desarrollo y medir el ahorro acumulado del piloto.')
    add_key_value_table(doc, [
        ('Horas ahorradas acumuladas', 'Pendiente: se obtiene a partir del piloto y los ciclos operativos validados.'),
        ('Horas totales de desarrollo', 'Pendiente: consolidar horas reales de levantamiento, desarrollo, pruebas e implementación.'),
        ('ROI estimado de tiempo', 'Pendiente de cálculo con las dos variables anteriores.'),
        ('ROI económico', 'Pendiente: requiere ahorro económico validado e inversión económica consolidada.'),
    ])

    add_heading(doc, '13. Resumen final de impacto')
    tbl = doc.add_table(rows=1, cols=2)
    style_table(tbl, ['Indicador', 'Resultado al 11 de septiembre de 2026'], [3.2, 4.8])
    for values in [
        ('Tiempo actual por ejecución', '12 horas por ciclo quincenal.'),
        ('Tiempo con la solución', 'Pendiente de piloto.'),
        ('Ejecuciones mensuales', '2 ciclos de referencia.'),
        ('Personas involucradas', '1 persona operativa.'),
        ('Horas actuales por mes', '24 horas.'),
        ('Horas ahorradas al mes', 'Pendiente de medir.'),
        ('Horas ahorradas al año', 'Pendiente de medir.'),
        ('Ahorro económico anual', 'Pendiente de medir; línea base de esfuerzo: $1.440/año.'),
        ('Inversión de tiempo de desarrollo', 'Pendiente de consolidar.'),
        ('ROI estimado de tiempo', 'Pendiente de cálculo con fórmula Quick Wins.'),
        ('ROI estimado económico', 'Pendiente de cálculo posterior al piloto.'),
    ]:
        row = tbl.add_row().cells
        set_cell_text(row[0], values[0], bold=True)
        shade(row[0], LIGHT_BLUE)
        set_cell_text(row[1], values[1])

    add_heading(doc, 'Siguientes acciones de validación', level=2)
    add_bullets(doc, [
        'Abrir SIAC y CAFI en la laptop Windows de Saskya y validar el mapa de controles de Ingreso de Procesos.',
        'Realizar un piloto supervisado con una fila autorizada y comprobar registro único, impresión y lectura de resultados.',
        'Probar actualización de Google Sheets con permisos autorizados y medir tiempos de proceso, supervisión e incidencias.',
    ])

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == '__main__':
    main()
