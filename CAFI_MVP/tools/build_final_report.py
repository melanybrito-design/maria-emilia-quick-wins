from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path('/Users/melanybrito/Desktop/CAFI_Documentación/Reporte de Avance Quickwins - Maria Emilia.docx')
OUTPUT = Path('/Users/melanybrito/Desktop/Maria Emilia (QuickWins) /CAFI_MVP/Reporte de Avance Quickwins - Maria Emilia Final.docx')
NAVY = '1F4E79'
LIGHT_BLUE = 'D9EAF7'
LIGHT_GRAY = 'E9EEF3'
BORDER = 'D9D9D9'


def set_cell_shading(cell, color):
    props = cell._tc.get_or_add_tcPr()
    shading = props.find(qn('w:shd'))
    if shading is None:
        shading = OxmlElement('w:shd')
        props.append(shading)
    shading.set(qn('w:fill'), color)


def set_cell_border(cell, color=BORDER):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for name in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = qn(f'w:{name}')
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f'w:{name}')
            borders.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:color'), color)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in('w:tcMar')
    if mar is None:
        mar = OxmlElement('w:tcMar')
        tc_pr.append(mar)
    for side, value in {'top': top, 'start': start, 'bottom': bottom, 'end': end}.items():
        node = mar.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}')
            mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    props = row._tr.get_or_add_trPr()
    item = OxmlElement('w:tblHeader')
    item.set(qn('w:val'), 'true')
    props.append(item)


def clear_body(document):
    body = document._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def style_run(run, bold=False, size=None, color=None, italic=False):
    run.bold = bold
    run.italic = italic
    run.font.name = 'Arial'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_paragraph(document, text='', style='Normal', bold_prefix=None, keep=False):
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.12
    paragraph.paragraph_format.keep_with_next = keep
    if bold_prefix and text.startswith(bold_prefix):
        style_run(paragraph.add_run(bold_prefix), bold=True)
        style_run(paragraph.add_run(text[len(bold_prefix):]))
    else:
        style_run(paragraph.add_run(text))
    return paragraph


def add_heading(document, text, level=1):
    paragraph = document.add_paragraph(style=f'Heading {level}')
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    style_run(run, bold=True, size=16 if level == 1 else 12, color='000000')
    return paragraph


def add_bullets(document, items):
    for item in items:
        paragraph = document.add_paragraph(style='List Bullet')
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.08
        style_run(paragraph.add_run(item))


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = 'Table Grid'
    head = table.rows[0]
    set_repeat_table_header(head)
    for idx, value in enumerate(headers):
        cell = head.cells[idx]
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(value)
        style_run(r, bold=True, size=9, color='FFFFFF')
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if widths:
            cell.width = Inches(widths[idx])
    for row_i, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cell = cells[idx]
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.03
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 and len(headers) > 2 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            style_run(r, bold=(idx == 0 and len(headers) == 2), size=8.6)
            if row_i % 2 == 1:
                set_cell_shading(cell, 'F6F8FA')
            if idx == 0 and len(headers) == 2:
                set_cell_shading(cell, LIGHT_BLUE)
            set_cell_border(cell)
            set_cell_margins(cell, 90, 100, 90, 100)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                cell.width = Inches(widths[idx])
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_metadata_table(document):
    rows = [
        ('Responsable operativa', 'Saskya Torres Guerrero, asistente administrativa de la Facultad de Comunicación'),
        ('Responsable de Quick Wins', 'María Emilia Aguirre Beltrán, impulsora y validadora del proyecto de automatización'),
        ('Apoyo y consultoría', 'Melany Brito, apoyo y consultoría principal para el levantamiento, documentación y desarrollo'),
        ('Área', 'Facultad de Comunicación, Universidad Espíritu Santo UEES'),
        ('Corte del informe', '11 de septiembre de 2026'),
    ]
    return add_table(document, ['Elemento', 'Información'], rows, [2.05, 4.55])


def main():
    document = Document(SOURCE)
    clear_body(document)
    for section in document.sections:
        section.top_margin = Inches(0.68)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.72)
        section.right_margin = Inches(0.72)

    title = document.add_paragraph(style='Title')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    style_run(title.add_run('REPORTE DE AVANCE QUICK WINS'), bold=True, size=22, color='000000')
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    style_run(subtitle.add_run('Automatización del ingreso de trámites desde Google Sheets o Excel hacia SIAC CAFI'), size=13, color='000000')
    add_paragraph(document, 'Este informe consolida el diagnóstico, el proceso actual, el avance técnico y el plan de validación del RPA para la Facultad de Comunicación. Al corte del 11 de septiembre de 2026, la conexión de lectura con SIAC se encuentra comprobada mediante UIA y Win32; no se han guardado trámites con el robot.', keep=True)
    add_metadata_table(document)

    add_heading(document, '1. Situación actual y punto de dolor')
    add_paragraph(document, 'Saskya Torres Guerrero realiza de forma manual e intermitente el ingreso de trámites de la Facultad de Comunicación. El proceso se ejecuta aproximadamente cada dos semanas y demanda 12 horas por ciclo. Como referencia operativa, equivale a 24 horas mensuales y 288 horas anuales de esfuerzo actual, asumiendo dos ciclos por mes. Estas cifras describen el tiempo base del proceso; todavía no constituyen ahorro validado.')
    add_paragraph(document, 'La carga exige trasladar información desde la agenda de Google Sheets o su exportación Excel hacia SIAC, módulo CAFI. Cada trámite requiere navegación, búsqueda del estudiante, ingreso de título y detalle, confirmación, impresión, lectura de resultados y actualización posterior de la agenda. La ejecución fragmentada no elimina el esfuerzo: extiende el ciclo y dificulta medir con precisión el tiempo por trámite.')
    add_heading(document, 'Puntos de dolor identificados', 2)
    add_bullets(document, [
        'Transcripción repetitiva de datos entre la agenda y CAFI, con riesgo de omisiones, cambios de título, pérdida de contexto o errores de digitación.',
        'Navegación recurrente por SIAC, CAFI, Procesos y Trámites Estudiante para iniciar cada registro.',
        'Dependencia de tiempos de espera de CAFI y de la verificación visual antes de guardar, imprimir o continuar.',
        'Necesidad de confirmar que el código ingresado resuelve el destinatario correcto antes de completar el trámite.',
        'Impresión individual y extracción manual posterior del código de trámite, GPA y créditos aprobados.',
        'Actualización posterior de la agenda principal, que hoy vuelve a requerir transferencia manual de los resultados.',
        'Riesgo de duplicar una carga si el sistema demora, se interrumpe o no queda claro si el guardado fue confirmado.',
    ])

    add_heading(document, '2. Proceso actual y resultado esperado')
    add_paragraph(document, 'La fuente de trabajo es una agenda organizada en Google Sheets. Para la automatización local, podrá usarse su exportación o un archivo Excel con la misma estructura. Los títulos, detalles y demás variables dependen de los datos vigentes de cada fila; no existe un catálogo fijo de títulos que el RPA deba imponer o modificar.')
    add_table(document, ['Paso', 'Actividad actual', 'Herramienta', 'Responsable', 'Resultado o control'], [
        ('1', 'Revisar la agenda y verificar que la fila tenga datos del estudiante, código, título y detalle.', 'Google Sheets o Excel', 'Saskya', 'Datos listos para cargar.'),
        ('2', 'Ingresar a SIAC con credenciales autorizadas y seleccionar CAFI.', 'SIAC', 'Saskya', 'Sesión disponible.'),
        ('3', 'Abrir Procesos, expandir la lista y elegir Trámites Estudiante > Crear nuevo proceso.', 'CAFI', 'Saskya', 'Formulario Ingreso de Procesos vacío.'),
        ('4', 'Ingresar el proceso, código del estudiante, título, detalle y prioridad según la agenda.', 'CAFI', 'Saskya', 'Campos completados y destinatario revisado.'),
        ('5', 'Guardar y esperar el mensaje de confirmación de la transacción.', 'CAFI', 'Saskya', 'Trámite registrado o incidencia identificada.'),
        ('6', 'Aceptar la confirmación e imprimir el trámite individual.', 'CAFI', 'Saskya', 'Reporte individual disponible.'),
        ('7', 'Leer código de trámite, GPA y créditos aprobados que devuelve el sistema.', 'CAFI', 'Saskya', 'Tres resultados disponibles.'),
        ('8', 'Actualizar la agenda principal con los tres resultados y preparar el siguiente trámite.', 'Google Sheets o Excel', 'Saskya', 'Agenda actualizada y trazabilidad por fila.'),
    ], [0.35, 2.35, 0.95, 0.75, 2.2])
    add_paragraph(document, 'Resultado esperado del proceso mejorado: cada fila de la agenda queda registrada en CAFI, impresa según el flujo vigente y actualizada en la agenda principal con el código de trámite, GPA y créditos aprobados. Si el sistema no confirma una acción, el caso debe quedar detenido para revisión y no reenviarse automáticamente.')

    add_heading(document, '3. Datos, herramientas y método de trabajo')
    add_table(document, ['Componente', 'Uso actual', 'Uso previsto en la solución'], [
        ('Agenda principal', 'Google Sheets con información de trámites.', 'Fuente de datos y destino de los tres resultados del trámite.'),
        ('Excel', 'Exportación o archivo de trabajo para la carga.', 'Validación local y respaldo operativo del lote.'),
        ('SIAC y CAFI', 'Aplicación de escritorio para registrar e imprimir trámites.', 'Sistema destino manejado en la misma laptop Windows de la usuaria.'),
        ('RPA local', 'No existe en la operación manual.', 'Ventana nativa que solicita acceso cuando corresponda, recibe el archivo y ejecuta el flujo controlado.'),
        ('Python, pywinauto, UIA y Win32', 'No participan en la operación manual.', 'Automatización de la interfaz Windows, lectura de controles y verificación de estado.'),
        ('Bitácora y copia de salida', 'Actualización manual de la agenda.', 'Registro por fila, estados, resultados leídos y soporte para conciliación.'),
    ], [1.3, 2.35, 2.95])
    add_paragraph(document, 'El método elegido es un RPA atendido y local. La asistente administra el inicio del proceso, revisa la agenda y mantiene la laptop disponible. El robot opera únicamente dentro de la sesión Windows autorizada, espera la respuesta de CAFI, verifica los campos antes de guardar y conserva una bitácora. No se plantea una API de SIAC porque no se dispone de evidencia de una interfaz programática institucional para este proceso.')

    add_heading(document, '4. Alcance de la solución')
    add_paragraph(document, 'El alcance del primer proceso automatizado inicia con una agenda validada y la sesión de SIAC disponible. El RPA toma los campos de cada fila, registra el trámite, espera la confirmación, ejecuta la impresión individual, recupera código de trámite, GPA y créditos aprobados, y prepara los resultados para actualizar la agenda principal. La actualización directa de Google Sheets es un requisito funcional del resultado final; la integración concreta con esa hoja se realizará después de definir el acceso autorizado y el mecanismo de actualización.')
    add_paragraph(document, 'El alcance no incluye aprobar solicitudes académicas, decidir el contenido de los trámites, cambiar títulos de la agenda, sustituir la revisión humana de inconsistencias ni automatizar procesos posteriores que no formen parte del ingreso e impresión de CAFI.')
    add_heading(document, 'Controles funcionales requeridos', 2)
    add_bullets(document, [
        'Mantener los ceros iniciales y el contenido literal de código, título y detalle.',
        'Aceptar títulos y datos variables provenientes de la agenda, sin sustituirlos por valores fijos.',
        'Comprobar que el estudiante o destinatario visible en CAFI corresponde a la fila antes de guardar.',
        'Registrar intención antes de guardar y marcar éxito solo cuando CAFI muestre la confirmación esperada.',
        'No reenviar automáticamente un trámite si la confirmación, impresión o lectura de resultados queda incompleta.',
        'Separar el estado de guardado del estado de impresión y de la actualización de resultados.',
    ])

    add_heading(document, '5. Responsables e involucrados')
    add_table(document, ['Rol', 'Responsabilidad'], [
        ('Saskya Torres Guerrero', 'Responsable operativa y usuaria principal. Ejecuta actualmente el proceso manual en la Facultad de Comunicación y participará en las pruebas operativas.'),
        ('María Emilia Aguirre Beltrán', 'Responsable de Quick Wins. Impulsa el aprendizaje y la implementación de la automatización, acompaña la validación del proyecto y articula el avance.'),
        ('Melany Brito', 'Apoyo y consultoría principal para el levantamiento, organización documental, definición de prompts y desarrollo técnico en Codex.'),
        ('Facultad de Comunicación', 'Área donde se aplica inicialmente el proceso y donde se valida su operación.'),
    ], [1.65, 4.95])

    add_heading(document, '6. Reuniones y avance del proyecto')
    add_table(document, ['Fecha', 'Enfoque', 'Resultados'], [
        ('8 de septiembre de 2026', 'Diagnóstico y levantamiento', 'Se identificó el proceso manual, la agenda como fuente de datos, el tiempo aproximado de 12 horas por ciclo quincenal, los puntos de dolor y el flujo hasta la impresión.'),
        ('11 de septiembre de 2026', 'Organización y desarrollo inicial', 'Se organizó la documentación, se estructuraron prompts, se inició el desarrollo en Codex y se preparó el trabajo en la laptop Windows para validar SIAC y CAFI.'),
    ], [1.25, 1.85, 3.5])
    add_heading(document, 'Evidencia técnica al corte', 2)
    add_bullets(document, [
        'Se creó una base de RPA local con interfaz de escritorio, validación de archivos, bitácora, control de estados y lógica de recuperación.',
        'La evidencia compartida registra conexión de lectura con SIAC por UIA y Win32, ambos con estado CONECTADO_LECTURA, además de inventarios de controles para continuar la configuración.',
        'La evidencia compartida registra 40 pruebas ejecutadas en Windows. Estas pruebas validan la base técnica; no equivalen a una carga real de trámites.',
        'No se han guardado trámites con el RPA. La carga permanece deshabilitada hasta verificar los controles del formulario Ingreso de Procesos y ejecutar el piloto autorizado.',
    ])

    add_heading(document, '7. Riesgos, dependencias y controles de mitigación')
    add_table(document, ['Riesgo o dependencia', 'Impacto', 'Control de mitigación'], [
        ('CAFÍ presenta tiempos de espera variables.', 'El robot puede avanzar antes de que el sistema responda.', 'Esperas basadas en la aparición o lectura de controles y detención ante tiempos excedidos.'),
        ('Campos de CAFI no verificados.', 'Podría escribirse en un campo incorrecto.', 'Mapeo de controles por UIA o Win32 en la laptop real; sin selectores inventados.'),
        ('Resultado incierto después de Guardar.', 'Riesgo de duplicar un trámite.', 'Persistir intención, detener el lote y exigir conciliación antes de un reintento.'),
        ('Discrepancia entre código y destinatario.', 'Registro asociado al estudiante equivocado.', 'Lectura de vuelta y comparación previa al guardado.'),
        ('Impresión o lectura incompleta.', 'La agenda podría quedar sin resultados o con datos inconsistentes.', 'Separar guardado, impresión, lectura y actualización; no volver a guardar una fila ya confirmada.'),
        ('Acceso a Google Sheets.', 'No se podría actualizar la agenda principal.', 'Definir cuenta, permisos, hoja y método de actualización antes de habilitar esa integración.'),
    ], [2.0, 1.75, 2.85])

    add_heading(document, '8. Medición de impacto')
    add_paragraph(document, 'La línea base confirmada es de 12 horas por ciclo quincenal, realizadas de manera intermitente por una persona. Para una referencia mensual de dos ciclos, el esfuerzo actual equivale a 24 horas. Con el costo de $5 por hora indicado por la plantilla Quick Wins, la capacidad asociada al proceso actual equivale a $120 mensuales y $1.440 anuales. Estas cifras no representan ahorro económico: son la base para medir el impacto una vez que el piloto determine el tiempo real con RPA y el tiempo humano residual.')
    add_table(document, ['Indicador', 'Situación al 11 de septiembre de 2026'], [
        ('Tiempo actual por ciclo', '12 horas, ejecutadas de manera intermitente.'),
        ('Frecuencia', 'Cada dos semanas; referencia de dos ciclos mensuales.'),
        ('Personas involucradas', 'Una persona operativa: Saskya Torres Guerrero.'),
        ('Horas actuales mensuales', '24 horas de referencia.'),
        ('Horas actuales anuales', '288 horas de referencia.'),
        ('Tiempo con RPA', 'Pendiente de piloto en Windows con SIAC y CAFI.'),
        ('Ahorro de tiempo y ROI', 'Pendientes de medición posterior al piloto.'),
    ], [2.25, 4.35])
    add_paragraph(document, 'Fórmula de medición posterior al piloto: horas ahorradas mensuales = horas actuales mensuales − horas mensuales con RPA. El tiempo con RPA debe incluir preparación de la agenda, supervisión, resolución de incidencias y conciliación de resultados, no solo la duración de la automatización.')

    add_heading(document, '9. Próximos pasos')
    add_bullets(document, [
        'Abrir SIAC y CAFI en la laptop Windows de Saskya, dejar disponible el formulario Ingreso de Procesos y validar el mapa de controles observado.',
        'Completar la configuración del formulario: acceso, navegación, código, destinatario, título, detalle, prioridad, confirmación, impresión, lectura de código de trámite, GPA y créditos aprobados.',
        'Realizar un piloto supervisado de una fila autorizada; comprobar que se registre una sola vez, se imprima y que los tres resultados correspondan a ese trámite.',
        'Probar dos filas consecutivas, una detención controlada y la conciliación de un caso incompleto antes de usar el proceso con un lote operativo.',
        'Definir el acceso autorizado a la agenda de Google Sheets y probar la actualización de resultados sin alterar campos fuente.',
        'Medir tiempo humano y tiempo total del piloto para calcular el ahorro y el ROI con evidencia operativa.',
    ])

    add_heading(document, '10. Conclusión')
    add_paragraph(document, 'El proyecto cuenta con un diagnóstico claro, una línea base de tiempo, una responsable operativa definida y una solución RPA en desarrollo. El principal valor esperado es eliminar la transferencia manual repetitiva entre la agenda y CAFI, preservar la trazabilidad y actualizar los resultados del trámite en la agenda principal. La siguiente etapa es técnica y controlada: validar los controles reales de CAFI en Windows, ejecutar un piloto supervisado y medir el desempeño antes de declarar ahorro o habilitar la operación por lotes.')
    add_heading(document, 'Criterios de aceptación del primer piloto', 2)
    add_table(document, ['Criterio', 'Evidencia requerida'], [
        ('Carga única', 'Una fila autorizada aparece una sola vez en CAFI y el mensaje de éxito corresponde a ese registro.'),
        ('Impresión y resultados', 'La impresión se completa y el código de trámite, GPA y créditos se identifican en el resultado del mismo trámite.'),
        ('Actualización de agenda', 'Los tres resultados se escriben en las columnas correctas sin modificar los datos fuente de la fila.'),
        ('Manejo de incidencia', 'Una detención o demora deja el caso identificado para conciliación y no genera un reintento automático.'),
        ('Medición', 'Se registra tiempo total, tiempo humano y número de trámites para comparar contra la línea base.'),
    ], [2.0, 4.6])

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.core_properties.title = 'Reporte de Avance Quick Wins'
    document.core_properties.subject = 'Automatización RPA del ingreso de trámites a CAFI'
    document.core_properties.author = 'Melany Brito'
    document.core_properties.comments = 'Informe final consolidado al 11 de septiembre de 2026.'
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == '__main__':
    main()
