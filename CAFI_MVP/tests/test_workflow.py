import io
import json
import threading
from pathlib import Path

import pytest
from openpyxl import load_workbook

from cafi_bot import SimulatedAdapter
from config import Rules
from excel_manager import InputError, digest, sample_book, validate
from journal import BusyError, Journal, worker_lock
from main import main
from runner import simulate

def modify(content, changes):
    wb = load_workbook(io.BytesIO(content))
    for address,value in changes.items():
        wb.active[address] = value
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()

@pytest.fixture
def book():
    return sample_book()

@pytest.fixture
def batch(tmp_path,book):
    journal = Journal(tmp_path/'data')
    run = journal.create(book,Rules())
    return journal,run['id']

def test_code_mask_and_preservation(book):
    wb = load_workbook(io.BytesIO(book))
    wb.active['A2'] = 17
    wb.active['A2'].number_format = '0000000000'
    stream = io.BytesIO(); wb.save(stream)
    result = validate(stream.getvalue())
    assert result['rows'][0]['code'] == '0000000017'
    assert result['rows'][1]['code'] == '9900000002'
    assert '\n' in result['rows'][0]['detail']
    assert 'revisión académica' in result['rows'][0]['detail']
    assert result['rows'][0]['priority'] == 'Media'
    assert not result['rows'][0]['issues']

@pytest.mark.parametrize('value',[None,17,1.5,-17,'00000000001','17','123456789x',True,'=17',99999999999999999])
def test_reject_invalid_codes(book,value):
    assert validate(modify(book,{'A2':value}))['rows'][0]['issues']

@pytest.mark.parametrize('changes',[{'A1':None},{'B1':'COD_ALUMNO'},{'H1':'TITULO'},{'H1':'DESCONOCIDO'},{'H2':'sin encabezado'}])
def test_schema_error(book,changes):
    with pytest.raises(InputError):
        validate(modify(book,changes))

def test_sheet_exact(book):
    with pytest.raises(InputError):
        validate(book,Rules(sheet='Hoja1'))

def test_formula_unknown_priority_and_empty_content(book):
    rows = validate(modify(book,{'H1':'PRIORIDAD','H2':'Alta','C3':'=1+2','D4':'   '}))['rows']
    assert all(row['issues'] for row in rows)

def test_duplicates_preserved(book):
    content = modify(book,{'A3':'9900000001','B3':'Estudiante de prueba 1'})
    rows = validate(content)['rows']
    assert len(rows)==3
    assert not rows[0]['issues'] and not rows[1]['issues']
    conflict = validate(modify(content,{'B3':'Otro nombre'}))['rows']
    assert conflict[0]['issues'] and conflict[1]['issues']
    duplicate = modify(content,{'C3':rows[0]['title'],'D3':rows[0]['detail']})
    assert validate(duplicate)['rows'][0]['issues']

def test_previous_results_blocked(book):
    assert validate(modify(book,{'E2':'123456'}))['rows'][0]['issues']

def test_empty_and_incomplete_rows(book):
    content = modify(book,{'A6':'9900000006'})
    diagnosis = validate(content)
    assert diagnosis['blank_rows']==1 and len(diagnosis['rows'])==4
    assert diagnosis['rows'][-1]['issues']

def test_default_dry_run_preserves_source_and_states(batch,book):
    journal,run_id=batch
    run=journal.get(run_id)
    assert run['counts']=={'PENDIENTE':3}
    assert (journal.root/run_id/'origen.xlsx').read_bytes()==book
    assert journal.create(book,Rules())['id']==run_id
    result=load_workbook(journal.export(run_id))
    source=load_workbook(io.BytesIO(book))
    for row in source.active:
        for cell in row:
            copied=result.active[cell.coordinate]
            assert copied.value==cell.value
            assert copied.number_format==cell.number_format
    assert result.active['J2'].value=='SIMULACIÓN / SIN CONEXIÓN CAFI'

def test_success_resume_does_not_resave(batch):
    journal,run_id=batch
    bot=SimulatedAdapter()
    assert simulate(journal,run_id,synthetic_confirmed=True,adapter=bot,limit=1)==3
    assert bot.saves==1
    assert simulate(journal,run_id,synthetic_confirmed=True,adapter=bot)==0
    assert bot.saves==3
    assert simulate(journal,run_id,synthetic_confirmed=True,adapter=bot)==0
    assert bot.saves==3
    run=journal.get(run_id)
    assert run['counts']=={'PROCESADO':3}
    assert sum(run['counts'].values())==run['total']

@pytest.mark.parametrize('scenario,state,code,saves',[
    ('timeout','REQUIERE_REVISION',3,1),('reject','ERROR',2,1),
    ('ok_failure','PROCESADO',3,1),('print_failure','PROCESADO',3,1),
    ('focus_loss','ERROR',3,0),('session_loss','ERROR',3,0),('readback_mismatch','ERROR',3,0)])
def test_failures_stop_and_preserve_outcome(batch,scenario,state,code,saves):
    journal,run_id=batch
    bot=SimulatedAdapter(scenario)
    assert simulate(journal,run_id,synthetic_confirmed=True,adapter=bot)==code
    assert bot.saves==saves
    run=journal.get(run_id)
    assert run['rows'][0]['state']==state
    assert run['rows'][1]['state']=='PENDIENTE'
    if state=='REQUIERE_REVISION':
        with pytest.raises(ValueError):
            simulate(journal,run_id,synthetic_confirmed=True,adapter=bot)
        assert bot.saves==1

def test_interrupted_intention_is_quarantined(batch):
    journal,run_id=batch
    journal.transition(run_id,2,'EN_PROCESO','INTENCION_GUARDAR','Intención de prueba.')
    with worker_lock(journal.root):
        journal.recover()
    assert journal.get(run_id)['rows'][0]['state']=='REQUIERE_REVISION'
    with pytest.raises(ValueError):
        simulate(journal,run_id,synthetic_confirmed=True)

def test_changed_source_blocks_resume(batch):
    journal,run_id=batch
    (journal.root/run_id/'origen.xlsx').write_bytes(b'cambio')
    with pytest.raises(ValueError,match='origen cambió'):
        simulate(journal,run_id,synthetic_confirmed=True)

def test_single_worker(batch):
    journal,run_id=batch
    with worker_lock(journal.root):
        with pytest.raises(BusyError):
            simulate(journal,run_id,synthetic_confirmed=True)

def test_stop_before_save(batch):
    journal,run_id=batch
    class StopAdapter(SimulatedAdapter):
        def prepare(self,row):
            journal.stop(run_id)
            return super().prepare(row)
    bot=StopAdapter()
    assert simulate(journal,run_id,synthetic_confirmed=True,adapter=bot)==3
    assert bot.saves==0
    assert journal.get(run_id)['rows'][0]['state']=='PENDIENTE'

def test_locked_output_after_success_does_not_resend(batch,monkeypatch):
    journal,run_id=batch
    original=journal.export
    calls=[0]
    def locked(run_id):
        calls[0]+=1
        if calls[0]>1:
            raise PermissionError('blocked')
        return original(run_id)
    monkeypatch.setattr(journal,'export',locked)
    bot=SimulatedAdapter()
    assert simulate(journal,run_id,synthetic_confirmed=True,adapter=bot)==3
    assert bot.saves==1
    assert journal.get(run_id)['rows'][0]['state']=='PROCESADO'
    monkeypatch.setattr(journal,'export',original)
    simulate(journal,run_id,synthetic_confirmed=True,adapter=bot)
    assert bot.saves==3

def test_failed_persist_after_save_leaves_uncertainty(batch,monkeypatch):
    journal,run_id=batch
    original=journal.transition
    def fail(run_id,row,state,*args,**kwargs):
        if state=='PROCESADO':
            raise OSError('disk full')
        return original(run_id,row,state,*args,**kwargs)
    monkeypatch.setattr(journal,'transition',fail)
    bot=SimulatedAdapter()
    assert simulate(journal,run_id,synthetic_confirmed=True,adapter=bot)==3
    assert bot.saves==1
    monkeypatch.setattr(journal,'transition',original)
    with worker_lock(journal.root):
        journal.recover()
    assert journal.get(run_id)['rows'][0]['state']=='REQUIERE_REVISION'

def test_invalid_rows_block_simulation(tmp_path):
    journal=Journal(tmp_path)
    run=journal.create(sample_book(True),Rules())
    with pytest.raises(ValueError):
        simulate(journal,run['id'],synthetic_confirmed=True)

def test_no_personal_content_in_support_report(batch):
    journal,run_id=batch
    simulate(journal,run_id,synthetic_confirmed=True)
    report=json.dumps(journal.support_report(run_id),ensure_ascii=False)
    for forbidden in ['9900000001','Estudiante de prueba','INCOMPLETO','revisión académica',str(journal.root)]:
        assert forbidden not in report

def test_cli_modes(tmp_path,book,capsys):
    path=tmp_path/'input.xlsx';path.write_bytes(book)
    flags=['--input',str(path),'--data-dir',str(tmp_path/'runs')]
    assert main(flags)==0
    assert 'PENDIENTE' in capsys.readouterr().out
    assert main(flags+['--execute'])==1
    assert main(flags+['--simulate'])==1
    assert main(flags+['--simulate','--synthetic-confirmed'])==0
    with pytest.raises(SystemExit):
        main(flags+['--dry-run','--execute'])
    with pytest.raises(SystemExit):
        main(flags+['--limit','0'])

def test_api_auth_and_flow(tmp_path):
    import re
    import urllib.error
    import urllib.request
    from server import make_server
    server=make_server(tmp_path,0)
    thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
    url=f'http://127.0.0.1:{server.server_port}'
    try:
        html=urllib.request.urlopen(url).read().decode()
        token=re.search(r'name="cafi-token" content="([^"]+)"',html).group(1)
        with pytest.raises(urllib.error.HTTPError) as error:
            urllib.request.urlopen(url+'/api/history')
        assert error.value.code==403
        with pytest.raises(urllib.error.HTTPError) as error:
            urllib.request.urlopen(urllib.request.Request(url+'/api/demo',data=b'{}',headers={'X-CAFI-Token':token,'Origin':'https://example.com'}))
        assert error.value.code==403
        request=urllib.request.Request(url+'/api/demo',data=b'{}',headers={'X-CAFI-Token':token})
        run=json.load(urllib.request.urlopen(request))
        assert run['total']==3
        request=urllib.request.Request(url+'/api/run?id='+run['id'],headers={'X-CAFI-Token':token})
        assert json.load(urllib.request.urlopen(request))['counts']=={'PENDIENTE':3}
    finally:
        server.shutdown();server.server_close();thread.join()
