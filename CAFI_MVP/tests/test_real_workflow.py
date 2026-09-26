import io
import json
import threading

import pytest
from openpyxl import load_workbook

from config import Rules
from excel_manager import sample_book
from journal import worker_lock
from real_workflow import RealJournal, execute_rows, start_workflow
from windows_rpa import RpaError, StopRequested, REQUIRED, empty_profile, validate_profile


class TestDriver:
    __test__ = False

    def __init__(self, journal, run_id, stop, failure=None):
        self.journal, self.run_id, self.stop = journal, run_id, stop
        self.failure = failure
        self.saves, self.prints, self.prepared = 0, 0, 0
        self.current = None

    def prepare(self, row):
        self.current = row
        self.prepared += 1
        assert row['code'].startswith('99')
        assert '\n' in row['detail']
        if self.failure == 'readback':
            raise RpaError('La lectura no coincide: título')
        if self.failure == 'stop_before_save':
            self.stop.set()

    def save(self):
        self.saves += 1
        persisted = next(r for r in self.journal.get(self.run_id)['rows'] if r['row'] == self.current['row'])
        assert persisted['stage'] == 'INTENCION_GUARDAR'
        if self.failure == 'timeout':
            raise RpaError('No se encontró la confirmación.')
        if self.failure == 'stop_during_save':
            self.stop.set()
            raise StopRequested('Detenido después de pulsar Guardar.')
        if self.failure == 'external_error':
            raise RuntimeError('SECRET-PASSWORD and arbitrary cell value')

    def open_print(self):
        assert any(r['state'] == 'PROCESADO' for r in self.journal.get(self.run_id)['rows'])
        if self.failure == 'ok':
            raise RpaError('No se encontró OK.')

    def print_report(self):
        assert self.journal.outputs(self.run_id)[self.current['row']]['phase'] == 'INTENCION_IMPRIMIR'
        self.prints += 1
        if self.failure == 'print':
            raise RpaError('No se abrió el reporte.')

    def results(self):
        if self.failure == 'results':
            raise RpaError('No se leyó el GPA.')
        return {'transaction_id': str(91000 + self.current['row']), 'gpa': '8.75', 'credits': '120'}

    def reset(self):
        if self.failure == 'reset':
            raise RpaError('No se limpió el formulario.')


def setup(tmp_path, failure=None):
    journal = RealJournal(tmp_path / 'operacion')
    content = sample_book()
    run = journal.create(content, Rules())
    stop = threading.Event()
    driver = TestDriver(journal, run['id'], stop, failure)
    return journal, run['id'], driver, stop, content


def test_real_pipeline_keeps_source_and_reads_result_values(tmp_path):
    journal, run_id, driver, stop, content = setup(tmp_path)
    with worker_lock(journal.root):
        assert execute_rows(journal, run_id, driver, stop) == 0
    assert driver.saves == driver.prints == 3
    assert journal.get(run_id)['status'] == 'COMPLETADO'
    assert all(r['phase'] == 'COMPLETO' for r in journal.outputs(run_id).values())
    assert (journal.root / run_id / 'origen.xlsx').read_bytes() == content
    output = journal.export(run_id)
    assert output.name == 'resultados_cafi.xlsx'
    wb = load_workbook(output)
    ws = wb.active
    assert ws['E2'].value == '91002'
    assert ws['F2'].value == '8.75'
    assert ws['G2'].value == '120'
    assert 'REFERENCIA_SIMULADA' not in [c.value for c in ws[1]]
    assert not (journal.root / run_id / 'resultados_simulacion.xlsx').exists()
    wb.close()
    with worker_lock(journal.root):
        assert execute_rows(journal, run_id, driver, stop) == 0
    assert driver.saves == driver.prints == 3  # Reabrir el mismo lote no reenvía.


@pytest.mark.parametrize('failure,state,saves,prints', [
    ('readback', 'ERROR', 0, 0), ('stop_before_save', 'ERROR', 0, 0),
    ('timeout', 'REQUIERE_REVISION', 1, 0), ('stop_during_save', 'REQUIERE_REVISION', 1, 0),
    ('external_error', 'REQUIERE_REVISION', 1, 0), ('ok', 'PROCESADO', 1, 0),
    ('print', 'PROCESADO', 1, 1), ('results', 'PROCESADO', 1, 1), ('reset', 'PROCESADO', 1, 1)])
def test_failures_never_repeat_submission(tmp_path, failure, state, saves, prints):
    journal, run_id, driver, stop, _ = setup(tmp_path, failure)
    with worker_lock(journal.root):
        assert execute_rows(journal, run_id, driver, stop) == 3
    result = journal.get(run_id)
    assert result['rows'][0]['state'] == state
    assert result['rows'][1]['state'] == 'PENDIENTE'
    assert driver.saves == saves and driver.prints == prints
    assert 'SECRET-PASSWORD' not in json.dumps(result)
    if state != 'ERROR':
        stop.clear()
        with worker_lock(journal.root), pytest.raises(RpaError):
            execute_rows(journal, run_id, driver, stop)
        assert driver.saves == saves


def test_crash_recovers_uncertain_save(tmp_path):
    journal, run_id, driver, stop, _ = setup(tmp_path)
    journal.transition(run_id, 2, 'EN_PROCESO', 'INTENCION_GUARDAR', 'Intento iniciado')
    with worker_lock(journal.root):
        journal.recover()
        with pytest.raises(RpaError):
            execute_rows(journal, run_id, driver, stop)
    assert driver.saves == 0


def test_duplicate_in_modified_workbook_is_blocked(tmp_path):
    journal, run_id, driver, stop, content = setup(tmp_path)
    with worker_lock(journal.root):
        execute_rows(journal, run_id, driver, stop)
    wb = load_workbook(io.BytesIO(content))
    wb.active['D4'] = 'Una solicitud nueva.'
    stream = io.BytesIO()
    wb.save(stream)
    wb.close()
    second = journal.create(stream.getvalue(), Rules())
    assert second['id'] != run_id
    with worker_lock(journal.root), pytest.raises(RpaError, match='intento previo'):
        execute_rows(journal, second['id'], driver, stop)
    assert driver.saves == 3


def test_output_failure_after_save_preserves_success(tmp_path, monkeypatch):
    journal, run_id, driver, stop, _ = setup(tmp_path)
    original = journal.export
    def locked_export(run_id):
        if journal.get(run_id)['rows'][0]['state'] == 'PROCESADO':
            raise PermissionError('Workbook locked')
        return original(run_id)
    monkeypatch.setattr(journal, 'export', locked_export)
    with worker_lock(journal.root):
        assert execute_rows(journal, run_id, driver, stop) == 3
    assert driver.saves == 1 and driver.prints == 0
    assert journal.get(run_id)['rows'][0]['state'] == 'PROCESADO'


def test_unpersisted_success_is_quarantined_on_recovery(tmp_path, monkeypatch):
    journal, run_id, driver, stop, _ = setup(tmp_path)
    original = journal.transition
    def transition(*args, **kwargs):
        if args[3] == 'EXITO':
            raise OSError('disk full')
        return original(*args, **kwargs)
    monkeypatch.setattr(journal, 'transition', transition)
    with worker_lock(journal.root):
        assert execute_rows(journal, run_id, driver, stop) == 3
        assert journal.get(run_id)['rows'][0]['state'] == 'EN_PROCESO'
        journal.recover()
    assert journal.get(run_id)['rows'][0]['state'] == 'REQUIERE_REVISION'
    assert driver.saves == 1


def test_real_output_does_not_convert_observed_text_to_formula(tmp_path):
    journal, run_id, _, _, _ = setup(tmp_path)
    journal.phase(run_id, 2, 'RESULTADOS_LEIDOS', {'gpa': '=HYPERLINK("example")', 'credits': '10'})
    wb = load_workbook(journal.export(run_id))
    assert wb.active['F2'].data_type == 's'
    wb.close()


def test_template_does_not_contain_invented_selectors():
    assert empty_profile()['controls'] == {}
    with pytest.raises(RpaError, match='Falta configurar'):
        validate_profile(empty_profile())


def test_missing_profile_never_touches_windows_or_stores_credentials(tmp_path, monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError('Windows must not be touched')
    monkeypatch.setattr('real_workflow.WindowsDriver', forbidden)
    with pytest.raises(RpaError, match='mapa'):
        start_workflow({'handle': 1}, 'file.xlsx', tmp_path / 'absent.json', tmp_path / 'data',
                       username='PRIVATE-USER', password='PRIVATE-PASSWORD')
    assert not (tmp_path / 'data').exists()


def test_limit_restarts_only_next_pending_row(tmp_path):
    journal, run_id, driver, stop, _ = setup(tmp_path)
    with worker_lock(journal.root):
        assert execute_rows(journal, run_id, driver, stop, limit=1) == 3
        assert driver.saves == 1
        assert execute_rows(journal, run_id, driver, stop) == 0
    assert driver.saves == 3


def test_human_confirmation_precedes_excel_result_update(tmp_path):
    journal, run_id, driver, stop, _ = setup(tmp_path)
    with worker_lock(journal.root):
        assert execute_rows(journal, run_id, driver, stop,
                            confirm_results=lambda row, result: False) == 3
    row = journal.get(run_id)['rows'][0]
    assert row['state'] == 'PROCESADO'
    assert journal.outputs(run_id)[row['row']]['phase'] == 'PENDIENTE_CONFIRMACION'
    output = journal.export(run_id)
    wb = load_workbook(output)
    assert wb.active['E2'].value in ('', None)
    wb.close()
